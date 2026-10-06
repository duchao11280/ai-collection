#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel Report Generator
A module for creating formatted Excel reports with configurable styles.
"""

import logging
from dataclasses import dataclass
from typing import Dict, List, Any, Optional, Union, Tuple
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


@dataclass
class ReportStyles:
    """Styles configuration for Excel report."""
    header_font: Font = Font(bold=True, color="FFFFFF")
    header_fill: PatternFill = PatternFill(
        start_color="000080", end_color="000080", fill_type="solid"
    )
    normal_font: Font = Font(name="Calibri", size=11)
    highlight_font: Font = Font(color="FF0000")
    center_alignment: Alignment = Alignment(horizontal='center', vertical='center')
    border: Border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )


class ExcelReportGenerator:
    """
    A class to generate formatted Excel reports.
    
    This class provides functionality to create professional Excel reports
    with consistent formatting, headers, data rows, and summary calculations.
    """
    
    def __init__(
        self, 
        headers: List[str],
        styles: Optional[ReportStyles] = None,
        percentage_columns: Optional[List[int]] = None,
        highlight_columns: Optional[List[int]] = None
    ):
        """
        Initialize the Excel report generator.
        
        Args:
            headers: List of column headers for the report
            styles: Optional custom styles configuration
            percentage_columns: Indices of columns to format as percentages
            highlight_columns: Indices of columns to highlight (e.g., with red font)
        """
        self.headers = headers
        self.styles = styles or ReportStyles()
        self.percentage_columns = percentage_columns or []
        self.highlight_columns = highlight_columns or []
        self.workbook: Optional[Workbook] = None
        self.worksheet: Optional[Worksheet] = None
    
    def _setup_workbook(self) -> None:
        """Create a new workbook and set up the active worksheet."""
        self.workbook = openpyxl.Workbook()
        self.worksheet = self.workbook.active
    
    def _write_headers(self) -> None:
        """Write and format the headers in the worksheet."""
        if not self.worksheet:
            raise ValueError("Worksheet is not initialized")
            
        for col_idx, header in enumerate(self.headers, 1):
            cell = self.worksheet.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.styles.header_font
            cell.fill = self.styles.header_fill
            cell.alignment = self.styles.center_alignment
            cell.border = self.styles.border
    
    def _format_cell(
        self, 
        row: int, 
        col: int, 
        value: Any,
        is_percentage: bool = False,
        highlight: bool = False
    ) -> None:
        """
        Format a cell with the appropriate styles.
        
        Args:
            row: Row index
            col: Column index
            value: Cell value
            is_percentage: Whether to format as percentage
            highlight: Whether to apply highlight formatting
        """
        if not self.worksheet:
            raise ValueError("Worksheet is not initialized")
            
        cell = self.worksheet.cell(row=row, column=col)
        
        # Set value with appropriate formatting
        if is_percentage and isinstance(value, (int, float)):
            cell.value = f"{value:.2f}%"
        else:
            cell.value = value
            
        # Apply styles
        cell.border = self.styles.border
        cell.alignment = self.styles.center_alignment
        
        if highlight:
            cell.font = self.styles.highlight_font
    
    def _write_data_rows(self, data: List[Dict[str, Any]]) -> None:
        """
        Write data rows to the worksheet.
        
        Args:
            data: List of dictionaries containing row data
        """
        if not self.worksheet:
            raise ValueError("Worksheet is not initialized")
            
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, header in enumerate(self.headers, 1):
                value = row_data.get(header, "")
                is_percentage = col_idx - 1 in self.percentage_columns
                highlight = col_idx - 1 in self.highlight_columns
                self._format_cell(row_idx, col_idx, value, is_percentage, highlight)
    
    def _write_totals_row(
        self, 
        data: List[Dict[str, Any]], 
        total_label: str = "TỔNG",
        numerical_columns: Optional[List[int]] = None
    ) -> None:
        """
        Write a summary row with totals for numerical columns.
        
        Args:
            data: List of dictionaries containing row data
            total_label: Label for the totals row
            numerical_columns: Indices of columns to sum (0-based)
        """
        if not self.worksheet or not data:
            return
            
        total_row_idx = len(data) + 2
        
        # Set the label for the totals row
        self._format_cell(total_row_idx, 1, total_label)
        
        # If numerical columns aren't specified, try to automatically detect them
        if not numerical_columns:
            numerical_columns = []
            for col_idx, header in enumerate(self.headers):
                if any(isinstance(row.get(header), (int, float)) for row in data):
                    numerical_columns.append(col_idx)
        
        # Calculate and write totals
        for col_idx in numerical_columns:
            header = self.headers[col_idx]
            total = sum(row.get(header, 0) for row in data if isinstance(row.get(header), (int, float)))
            
            is_percentage = col_idx in self.percentage_columns
            highlight = col_idx in self.highlight_columns
            
            self._format_cell(
                total_row_idx, 
                col_idx + 1, 
                total, 
                is_percentage, 
                highlight
            )
            
        # Format any remaining cells in the total row
        for col_idx in range(1, len(self.headers) + 1):
            if self.worksheet.cell(row=total_row_idx, column=col_idx).value is None:
                self._format_cell(total_row_idx, col_idx, "")
    
    def _calculate_special_totals(
        self, 
        data: List[Dict[str, Any]],
        total_row_idx: int,
        calculations: List[Tuple[int, str, List[str]]]
    ) -> None:
        """
        Calculate special totals that require custom logic.
        
        Args:
            data: List of dictionaries containing row data
            total_row_idx: Index of the totals row
            calculations: List of tuples containing (column_index, formula_type, source_columns)
        """
        if not self.worksheet:
            return
            
        for col_idx, formula_type, source_columns in calculations:
            if formula_type == "percentage":
                # Calculate percentage (e.g., for "Tỷ lệ NG")
                numerator = sum(row.get(source_columns[0], 0) for row in data)
                denominator = sum(row.get(source_columns[1], 0) for row in data)
                
                if denominator > 0:
                    value = (numerator / denominator) * 100
                else:
                    value = 0
                    
                is_percentage = True
                highlight = col_idx - 1 in self.highlight_columns
                self._format_cell(total_row_idx, col_idx, value, is_percentage, highlight)
    
    def _adjust_column_widths(self) -> None:
        """Automatically adjust column widths for better readability."""
        if not self.worksheet:
            return
            
        for col_idx, header in enumerate(self.headers, 1):
            # Find the maximum width needed
            max_length = len(str(header))
            for row_idx in range(2, self.worksheet.max_row + 1):
                cell_value = self.worksheet.cell(row=row_idx, column=col_idx).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            
            # Set column width with some padding
            column = get_column_letter(col_idx)
            self.worksheet.column_dimensions[column].width = max_length + 4
    
    def generate_report(
        self, 
        data: List[Dict[str, Any]], 
        output_filename: str, 
        include_totals: bool = True,
        auto_adjust_width: bool = True,
        special_calculations: Optional[List[Tuple[int, str, List[str]]]] = None
    ) -> str:
        """
        Generate a formatted Excel report.
        
        Args:
            data: List of dictionaries containing row data
            output_filename: Name of the output Excel file
            include_totals: Whether to include a totals row
            auto_adjust_width: Whether to adjust column widths automatically
            special_calculations: List of special calculations for the totals row
            
        Returns:
            Path to the generated Excel file
        """
        try:
            self._setup_workbook()
            self._write_headers()
            self._write_data_rows(data)
            
            if include_totals:
                numerical_columns = [i for i, header in enumerate(self.headers) 
                                   if any(isinstance(row.get(header), (int, float)) for row in data)]
                
                # Write regular totals
                total_row_idx = len(data) + 2
                self._write_totals_row(data, numerical_columns=numerical_columns)
                
                # Apply any special calculations
                if special_calculations:
                    self._calculate_special_totals(data, total_row_idx, special_calculations)
            
            if auto_adjust_width:
                self._adjust_column_widths()
            
            # Save the workbook
            if self.workbook:
                self.workbook.save(output_filename)
                logger.info(f"Excel report generated: {output_filename}")
                return output_filename
                
        except Exception as e:
            logger.error(f"Error generating Excel report: {str(e)}")
            raise
            
        return ""


def create_quality_report(
    data: List[Dict[str, Any]], 
    output_filename: str = "quality_report.xlsx"
) -> str:
    """
    Create a quality inspection report.
    
    Args:
        data: List of dictionaries containing quality inspection data
        output_filename: Name of the output Excel file
        
    Returns:
        Path to the generated Excel file
    """
    # Define headers
    headers = ["STT", "Ngày kiểm tra", "VISION", "Ca làm việc", 
               "Số lượng KT", "OK", "NG", "Tỷ lệ NG"]
    
    # Configure special formatting
    percentage_columns = [7]  # "Tỷ lệ NG" column index (0-based)
    highlight_columns = [7]   # "Tỷ lệ NG" column index (0-based)
    
    # Create report generator
    generator = ExcelReportGenerator(
        headers=headers,
        percentage_columns=percentage_columns,
        highlight_columns=highlight_columns
    )
    
    # Define special calculations for the totals row
    # For "Tỷ lệ NG" total, we need to calculate (total NG / total Số lượng KT) * 100
    special_calculations = [
        (8, "percentage", ["NG", "Số lượng KT"])
    ]
    
    # Generate the report
    return generator.generate_report(
        data=data,
        output_filename=output_filename,
        include_totals=True,
        auto_adjust_width=True,
        special_calculations=special_calculations
    )


def main():
    """Main function to demonstrate the report generation."""
    # Sample data
    sample_data = [
        {
            "STT": 1,
            "Ngày kiểm tra": "2025/03/01",
            "VISION": "N",
            "Ca làm việc": 100000,
            "Số lượng KT": 100000,
            "OK": 99999,
            "NG": 1,
            "Tỷ lệ NG": 0.00
        },
        {
            "STT": 1,
            "Ngày kiểm tra": "2025/03/01",
            "VISION": "D",
            "Ca làm việc": 100000,
            "Số lượng KT": 100000,
            "OK": 99990,
            "NG": 10,
            "Tỷ lệ NG": 0.01
        },
        {
            "STT": 2,
            "Ngày kiểm tra": "2025/03/02",
            "VISION": "N",
            "Ca làm việc": 180000,
            "Số lượng KT": 180000,
            "OK": 99990,
            "NG": 80010,
            "Tỷ lệ NG": 44.45
        },
        {
            "STT": 3,
            "Ngày kiểm tra": "2025/03/03",
            "VISION": "D",
            "Ca làm việc": 160000,
            "Số lượng KT": 160000,
            "OK": 60000,
            "NG": 100000,
            "Tỷ lệ NG": 62.50
        },
        {
            "STT": 3,
            "Ngày kiểm tra": "2025/03/03",
            "VISION": "N",
            "Ca làm việc": 170000,
            "Số lượng KT": 170000,
            "OK": 60000,
            "NG": 110000,
            "Tỷ lệ NG": 64.71
        },
        {
            "STT": 4,
            "Ngày kiểm tra": "2025/03/04",
            "VISION": "D",
            "Ca làm việc": 170000,
            "Số lượng KT": 170000,
            "OK": 60000,
            "NG": 110000,
            "Tỷ lệ NG": 64.71
        }
    ]
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = create_quality_report(sample_data, f"quality_report_{timestamp}.xlsx")
    print(f"Excel report generated: {output_file}")


if __name__ == "__main__":
    main()